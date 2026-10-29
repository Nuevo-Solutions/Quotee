import json
from io import BytesIO
from types import SimpleNamespace

import httpx
import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.catalog.importer import parse_catalog_file
from app.db.base import Base
from app.db.session import get_db_session
from app.llm.client import LLMTokenUsage
from app.main import app
from app.schemas.catalog import CatalogColumnMapping

MAPPING = {
    "sku": 0,
    "name": 1,
    "unit": 2,
    "price": 3,
    "currency": 4,
    "vat_rate": 5,
    "aliases": 6,
}


@pytest_asyncio.fixture
async def catalog_database(tmp_path):
    database_path = tmp_path / "catalog-test.db"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(engine, expire_on_commit=False)

    async def override_session():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db_session] = override_session
    try:
        yield
    finally:
        app.dependency_overrides.pop(get_db_session, None)
        await engine.dispose()


def make_xlsx_bytes() -> bytes:
    workbook = Workbook()
    products = workbook.active
    products.title = "Artikli"
    products.append(["SKU", "Proizvod", "JM", "Cena", "Valuta", "PDV", "Alias"])
    products.append(["X-1", "Kabl", "m", 12.5, "RSD", 20, "Žica|Kabl"])
    workbook.create_sheet("Drugi list").append(["Nepovezano"])

    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def make_catalog_with_title_rows() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Katalog"
    sheet.append(["Demo katalog proizvoda"])
    sheet.append([])
    sheet.append(
        ["R. br.", "Šifra", "Naziv proizvoda", "Jedinica", "Cena", "Valuta", "PDV", "Aktivan"]
    )
    sheet.append([1, "A-1", "Kabl PP-Y", "m", 125.5, "RSD", 20, True])
    sheet.append([2, "A-2", "Ukinuti proizvod", "kom", 10, "RSD", 20, False])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def make_catalog_after_cover_sheet() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Naslovna"
    workbook.active.append(["Uvodna strana"])
    catalog = workbook.create_sheet("Katalog")
    catalog.append(["SKU", "Proizvod", "JM"])
    catalog.append(["A-1", "Kabl", "m"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def test_catalog_auto_detects_header_after_title_and_preserves_active_flag() -> None:
    parsed = parse_catalog_file(
        filename="katalog.xlsx",
        file_bytes=make_catalog_with_title_rows(),
    )

    assert parsed.sheet_name == "Katalog"
    assert [product.source_row for product in parsed.products] == [4, 5]
    assert [product.sku for product in parsed.products] == ["A-1", "A-2"]
    assert [product.active for product in parsed.products] == [True, False]


def test_catalog_auto_selects_data_sheet_after_cover_sheet() -> None:
    parsed = parse_catalog_file(
        filename="katalog-sa-naslovnom.xlsx",
        file_bytes=make_catalog_after_cover_sheet(),
    )

    assert parsed.sheet_name == "Katalog"
    assert [product.sku for product in parsed.products] == ["A-1"]


@pytest.mark.asyncio
async def test_catalog_api_imports_without_mapping(catalog_database) -> None:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/catalog/imports",
            files={
                "file": (
                    "automatski-katalog.xlsx",
                    make_catalog_with_title_rows(),
                    "application/octet-stream",
                )
            },
        )

    assert response.status_code == 201
    assert response.json()["sheet_name"] == "Katalog"
    assert response.json()["products_imported"] == 2


@pytest.mark.asyncio
async def test_csv_catalog_mapping_imports_versioned_products_and_aliases(catalog_database) -> None:
    csv_content = (
        "SKU;Proizvod;JM;Cena;Valuta;PDV;Alternativni naziv\n"
        "CAB-001;Kabl PP-Y;m;125,50;RSD;20;Žica|Kabl\n"
        "BRK-001;Osigurač;kom;42,00;RSD;20;Automat\n"
    ).encode()
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        first_import = await client.post(
            "/api/v1/catalog/imports",
            data={"mapping": json.dumps(MAPPING)},
            files={"file": ("katalog.csv", csv_content, "text/csv")},
        )
        second_import = await client.post(
            "/api/v1/catalog/imports",
            data={"mapping": json.dumps(MAPPING)},
            files={"file": ("katalog.csv", csv_content, "text/csv")},
        )
        imports = await client.get("/api/v1/catalog/imports")
        products = await client.get(
            "/api/v1/catalog/products",
            params={"catalog_import_id": first_import.json()["id"]},
        )

    assert first_import.status_code == 201
    assert first_import.json()["version"] == 1
    assert first_import.json()["products_imported"] == 2
    assert first_import.json()["aliases_imported"] == 3
    assert second_import.status_code == 201
    assert second_import.json()["version"] == 2
    assert [item["version"] for item in imports.json()] == [2, 1]
    products_by_sku = {product["sku"]: product for product in products.json()}
    assert products_by_sku["CAB-001"]["price"] == "125.5000"
    assert products_by_sku["CAB-001"]["aliases"] == ["Kabl", "Žica"]


@pytest.mark.asyncio
async def test_xlsx_catalog_uses_selected_worksheet(catalog_database) -> None:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/catalog/imports",
            data={"mapping": json.dumps(MAPPING), "sheet_name": "Artikli"},
            files={"file": ("katalog.xlsx", make_xlsx_bytes(), "application/octet-stream")},
        )

    assert response.status_code == 201
    assert response.json()["sheet_name"] == "Artikli"
    assert response.json()["products_imported"] == 1


@pytest.mark.asyncio
async def test_invalid_catalog_rows_are_rejected_without_creating_import(catalog_database) -> None:
    csv_content = (
        b"SKU,Proizvod,JM,Cena,Valuta,PDV,Alias\n"
        b"DUP-1,Kabl,m,12,RSD,20,Wire\n"
        b"DUP-1,Drugi kabl,m,13,RSD,20,Cable\n"
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/catalog/imports",
            data={"mapping": json.dumps(MAPPING)},
            files={"file": ("duplicates.csv", csv_content, "text/csv")},
        )
        imports = await client.get("/api/v1/catalog/imports")

    assert response.status_code == 422
    assert "dupliran" in response.json()["detail"]["errors"][0]
    assert imports.json() == []


@pytest.mark.asyncio
async def test_catalog_import_rejects_missing_required_mapping(catalog_database) -> None:
    invalid_mapping = {**MAPPING, "unit": None}
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/catalog/imports",
            data={"mapping": json.dumps(invalid_mapping)},
            files={"file": ("catalog.csv", b"sku,name,unit\nA-1,Cable,m\n", "text/csv")},
        )

    assert response.status_code == 422
    assert response.json()["detail"]["message"] == "Mapiranje kolona nije validno."


@pytest.mark.asyncio
async def test_catalog_import_uses_bounded_llm_fallback_without_mapping(
    catalog_database,
    monkeypatch,
) -> None:
    from app.api.v1.endpoints import catalog as catalog_endpoint

    class FakeLLM:
        def __init__(self) -> None:
            self.headers = None
            self.sample_rows = None

        async def suggest_catalog_column_mapping(self, *, headers, sample_rows):
            self.headers = headers
            self.sample_rows = sample_rows
            return (
                CatalogColumnMapping(sku=0, name=1, unit=2, price=3),
                LLMTokenUsage(input_tokens=50, output_tokens=20),
            )

        async def close(self):
            return None

    fake_llm = FakeLLM()
    monkeypatch.setattr(
        catalog_endpoint,
        "get_settings",
        lambda: SimpleNamespace(
            max_document_size_bytes=1024 * 1024,
            llm_job_max_tokens=30000,
        ),
    )
    monkeypatch.setattr(
        catalog_endpoint,
        "create_llm_client",
        lambda *_args, **_kwargs: fake_llm,
    )
    catalog_csv = b"Artikel ID,Artikelbeschreibung,Verkaufseinheit,Preis\nA-1,Blue Cable,m,12.00\n"

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/catalog/imports",
            files={"file": ("unmapped.csv", catalog_csv, "text/csv")},
        )

    assert response.status_code == 201
    assert response.json()["products_imported"] == 1
    assert fake_llm.headers == ["Artikel ID", "Artikelbeschreibung", "Verkaufseinheit", "Preis"]
    assert fake_llm.sample_rows == [["", "", "", ""]]
