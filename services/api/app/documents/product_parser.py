import csv
import io
import re
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.documents.layout import (
    HeaderCandidate,
    best_preview_row,
    find_header_candidates,
    normalize_header,
)


class UnsupportedProductDocumentError(ValueError):
    pass


class ProductColumnMappingRequiredError(UnsupportedProductDocumentError):
    def __init__(
        self,
        headers: list[str],
        sample_rows: list[list[str]],
        *,
        header_row: int = 0,
        sheet_name: str | None = None,
    ) -> None:
        super().__init__("Potrebno je potvrditi mapiranje kolona.")
        self.headers = headers
        self.sample_rows = sample_rows
        self.header_row = header_row
        self.sheet_name = sheet_name


@dataclass(frozen=True)
class ParsedProductLine:
    position_code: str | None
    raw_description: str
    quantity: Decimal | None
    unit: str | None
    source_sku: str | None
    source_ref: dict[str, object]
    flags: list[str]


HEADER_ALIASES = {
    "position": {
        "position",
        "pozicija",
        "poz",
        "red",
        "rbr",
        "r. br.",
        "redni broj",
        "line",
    },
    "sku": {
        "sku",
        "sifra",
        "sifraartikla",
        "sifra artikla",
        "sifra proizvoda",
        "artikalsifra",
        "code",
        "productcode",
        "kod",
    },
    "description": {
        "description",
        "opis",
        "opis artikla",
        "naziv",
        "naziv proizvoda",
        "proizvod",
        "name",
        "artikal",
        "artikl",
        "item",
    },
    "quantity": {"quantity", "kolicina", "qty", "kol", "kol."},
    "unit": {
        "unit",
        "jedinica",
        "jedinica mere",
        "jedmere",
        "jed. mere",
        "jm",
        "j.m.",
        "um",
        "measure",
    },
}
QUANTITY_PATTERN = re.compile(r"^\s*[-+]?\d+(?:[.,]\d+)?\s*$")
SKU_PATTERN = re.compile(r"^[\w./-]{1,120}$", re.UNICODE)
MAX_PRODUCT_POSITIONS = 100
MAX_MAPPABLE_COLUMNS = 10


def _column_map(header: list[object]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(header):
        normalized = normalize_header(cell)
        for field, aliases in HEADER_ALIASES.items():
            if normalized in aliases and field not in columns:
                columns[field] = index
    return columns


def _read_csv(file_bytes: bytes) -> list[list[object]]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise UnsupportedProductDocumentError("CSV mora biti kodiran kao UTF-8.") from error

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        return [
            list(row) for row in csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
        ]
    except csv.Error as error:
        raise UnsupportedProductDocumentError("CSV struktura nije validna.") from error


def _read_xlsx_sheets(file_bytes: bytes) -> list[tuple[str, list[list[object]]]]:
    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
        raise UnsupportedProductDocumentError("XLSX fajl nije validna Excel radna sveska.")
    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError) as error:
        raise UnsupportedProductDocumentError("XLSX fajl nije moguće otvoriti.") from error

    try:
        return [
            (worksheet.title, [list(row) for row in worksheet.iter_rows(values_only=True)])
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _quantity_value(value: object, row_number: int) -> tuple[Decimal | None, list[str]]:
    if value is None or str(value).strip() == "":
        return None, ["MISSING_QUANTITY"]

    raw_value = str(value).strip().replace("\u00a0", " ")
    raw_value = raw_value.replace(" ", "")
    if "," in raw_value and "." in raw_value:
        if raw_value.rfind(",") > raw_value.rfind("."):
            raw_value = raw_value.replace(".", "").replace(",", ".")
        else:
            raw_value = raw_value.replace(",", "")
    elif raw_value.count(",") > 1:
        parts = raw_value.split(",")
        if all(len(part) == 3 for part in parts[1:]):
            raw_value = "".join(parts)
        else:
            return None, ["INVALID_QUANTITY"]
    elif raw_value.count(".") > 1:
        parts = raw_value.split(".")
        if all(len(part) == 3 for part in parts[1:]):
            raw_value = "".join(parts)
        else:
            return None, ["INVALID_QUANTITY"]
    elif "," in raw_value:
        decimal_places = len(raw_value.rsplit(",", maxsplit=1)[1])
        if decimal_places == 3:
            return None, ["INVALID_QUANTITY"]
        raw_value = raw_value.replace(",", ".")
    elif "." in raw_value and len(raw_value.rsplit(".", maxsplit=1)[1]) == 3:
        return None, ["INVALID_QUANTITY"]

    if not QUANTITY_PATTERN.fullmatch(raw_value):
        return None, ["INVALID_QUANTITY"]

    try:
        quantity = Decimal(raw_value.replace(",", "."))
    except InvalidOperation:
        return None, ["INVALID_QUANTITY"]
    if not quantity.is_finite() or quantity <= 0:
        return quantity, ["INVALID_QUANTITY"]
    return quantity, []


def _cell_value(row: list[object], columns: dict[str, int], field: str) -> object | None:
    column_index = columns.get(field)
    return row[column_index] if column_index is not None and column_index < len(row) else None


def _sku_value(value: object | None) -> str | None:
    if value is None:
        return None
    sku = str(value).strip()
    if not sku or any(character.isspace() for character in sku) or not SKU_PATTERN.fullmatch(sku):
        return None
    return sku


def _is_product_row(row: list[object], columns: dict[str, int]) -> bool:
    sku = _sku_value(_cell_value(row, columns, "sku"))
    description = str(_cell_value(row, columns, "description") or "").strip()
    has_row_signal = any(
        _cell_value(row, columns, field) is not None
        and bool(str(_cell_value(row, columns, field)).strip())
        for field in ("position", "quantity", "unit")
    )
    has_requested_values = any(
        _cell_value(row, columns, field) is not None
        and bool(str(_cell_value(row, columns, field)).strip())
        for field in ("quantity", "unit")
    )
    return bool(
        sku or (description and has_row_signal) or (has_row_signal and has_requested_values)
    )


def _find_product_headers(rows: list[list[object]]) -> list[HeaderCandidate]:
    return find_header_candidates(
        rows,
        aliases=HEADER_ALIASES,
        identity_fields={"sku", "description"},
        minimum_fields=2,
        row_validator=_is_product_row,
    )


def _preview_headers(rows: list[list[object]], row_index: int) -> list[str]:
    headers = [
        "" if value is None else str(value).strip()
        for value in rows[row_index][:MAX_MAPPABLE_COLUMNS]
    ]
    headers.extend([""] * (MAX_MAPPABLE_COLUMNS - len(headers)))
    return headers


def _parse_table_rows(
    *,
    rows: list[list[object]],
    sheet_name: str | None,
    document_id: UUID,
    header: HeaderCandidate,
    other_header_rows: set[int],
    explicit_mapping: dict[str, int | None] | None = None,
) -> list[ParsedProductLine]:
    columns = explicit_mapping or header.columns
    parsed_lines: list[ParsedProductLine] = []
    empty_streak = 0
    noise_streak = 0

    for row_index in range(header.row_index + 1, len(rows)):
        if row_index in other_header_rows:
            break
        row = rows[row_index]
        row_number = row_index + 1
        if not any(cell is not None and str(cell).strip() for cell in row):
            empty_streak += 1
            noise_streak += 1
            if empty_streak >= 3 or noise_streak >= 3:
                break
            continue
        empty_streak = 0
        if not _is_product_row(row, columns):
            noise_streak += 1
            if noise_streak >= 3:
                break
            continue
        noise_streak = 0
        if len(parsed_lines) >= MAX_PRODUCT_POSITIONS:
            raise UnsupportedProductDocumentError(
                f"Product dokument prelazi lokalni limit od {MAX_PRODUCT_POSITIONS} pozicija."
            )

        source_sku = _sku_value(_cell_value(row, columns, "sku"))
        description_value = _cell_value(row, columns, "description")
        description = str(description_value).strip() if description_value is not None else ""
        if not description:
            description = source_sku or ""

        quantity, flags = _quantity_value(_cell_value(row, columns, "quantity"), row_number)
        if not description:
            flags.append("MISSING_DESCRIPTION")
        unit_value = _cell_value(row, columns, "unit")
        unit = (
            str(unit_value).strip() if unit_value is not None and str(unit_value).strip() else None
        )
        if unit is None:
            flags.append("MISSING_UNIT")

        position_value = _cell_value(row, columns, "position")
        position_code = (
            str(position_value).strip()
            if position_value is not None and str(position_value).strip()
            else str(len(parsed_lines) + 1)
        )
        source_ref: dict[str, object] = {
            "document_id": str(document_id),
            "row": row_number,
            "columns": dict(columns),
        }
        if sheet_name is not None:
            source_ref["sheet"] = sheet_name
        if source_sku:
            source_ref["sku"] = source_sku

        parsed_lines.append(
            ParsedProductLine(
                position_code=position_code,
                raw_description=description,
                quantity=quantity,
                unit=unit,
                source_sku=source_sku,
                source_ref=source_ref,
                flags=flags,
            )
        )
    return parsed_lines


def parse_product_document(
    *,
    document_id: UUID,
    filename: str,
    file_bytes: bytes,
    column_mapping: dict[str, int | None] | None = None,
    sheet_name: str | None = None,
    header_row: int | None = None,
) -> tuple[list[ParsedProductLine], str | None]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        sheets = [(None, _read_csv(file_bytes))]
    elif extension == ".xlsx":
        sheets = _read_xlsx_sheets(file_bytes)
        if sheet_name is not None:
            sheets = [sheet for sheet in sheets if sheet[0] == sheet_name]
            if not sheets:
                raise UnsupportedProductDocumentError("Izabrani radni list više nije dostupan.")
    else:
        raise UnsupportedProductDocumentError("Product obrada trenutno podržava CSV i XLSX.")

    if not any(rows for _, rows in sheets):
        raise UnsupportedProductDocumentError("Dokument nema redove za obradu.")

    if column_mapping is not None:
        if any(
            field not in HEADER_ALIASES
            or index is not None
            and not 0 <= index < MAX_MAPPABLE_COLUMNS
            for field, index in column_mapping.items()
        ):
            raise UnsupportedProductDocumentError("Mapiranje kolona nije validno.")

    lines: list[ParsedProductLine] = []
    preview: tuple[str | None, list[list[object]], int] | None = None
    for sheet_name, rows in sheets:
        if not rows:
            continue
        headers = _find_product_headers(rows)
        if column_mapping is not None:
            selected_header_row = (
                header_row
                if header_row is not None
                else headers[0].row_index
                if headers
                else best_preview_row(rows)
            )
            headers = [
                HeaderCandidate(
                    row_index=selected_header_row,
                    columns={
                        field: index for field, index in column_mapping.items() if index is not None
                    },
                    score=0,
                    supporting_rows=0,
                )
            ]
        if not headers:
            row_index = best_preview_row(rows)
            row_score = sum(
                value is not None and bool(str(value).strip()) for value in rows[row_index]
            )
            if preview is None or row_score > sum(
                value is not None and bool(str(value).strip()) for value in preview[1][preview[2]]
            ):
                preview = (sheet_name, rows, row_index)
            continue

        header_rows = {header.row_index for header in headers}
        for header in headers:
            lines.extend(
                _parse_table_rows(
                    rows=rows,
                    sheet_name=sheet_name,
                    document_id=document_id,
                    header=header,
                    other_header_rows=header_rows - {header.row_index},
                    explicit_mapping=column_mapping,
                )
            )

    if not lines:
        if preview is None:
            raise UnsupportedProductDocumentError("Nijedna stavka nije pronađena u dokumentu.")
        sheet_name, rows, header_row = preview
        raise ProductColumnMappingRequiredError(
            headers=_preview_headers(rows, header_row),
            sample_rows=[
                [
                    "" if value is None else str(value).strip()
                    for value in row[:MAX_MAPPABLE_COLUMNS]
                ]
                for row in rows[header_row + 1 : header_row + 1 + 5]
            ],
            header_row=header_row,
            sheet_name=sheet_name,
        )
    result_sheet = lines[0].source_ref.get("sheet") if len(sheets) == 1 else None
    return lines, str(result_sheet) if result_sheet is not None else None
