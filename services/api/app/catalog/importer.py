import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.documents.layout import HeaderCandidate, find_header_candidates, normalize_header
from app.schemas.catalog import CatalogColumnMapping, CatalogProductInput

MAX_CATALOG_PRODUCTS = 3000
MAX_ROW_ERRORS = 100
ALIAS_SEPARATOR = re.compile(r"[;|]")
CATALOG_HEADER_ALIASES = {
    "sku": {
        "sku",
        "sifra",
        "sifra artikla",
        "sifra proizvoda",
        "artikalsifra",
        "code",
        "product code",
        "artikalkod",
        "kod",
    },
    "name": {
        "name",
        "naziv",
        "naziv proizvoda",
        "proizvod",
        "product",
        "product name",
        "opis",
        "opis artikla",
        "description",
        "artikal",
        "artikl",
    },
    "unit": {"unit", "jm", "j.m.", "jedinica", "jedinica mere", "jed. mere", "mera"},
    "price": {"price", "cena", "cena bez pdv", "jedinicna cena", "unit price"},
    "currency": {"currency", "valuta", "curr"},
    "vat_rate": {"vat rate", "vat", "pdv", "pdv stopa", "stopa pdv"},
    "aliases": {"alias", "aliases", "alternativni naziv", "alternativni nazivi"},
    "active": {"active", "aktivan", "aktivno", "status"},
}
ACTIVE_TRUE = {"1", "true", "yes", "y", "da", "aktivan", "aktivno"}
ACTIVE_FALSE = {"0", "false", "no", "n", "ne", "neaktivan", "neaktivno"}


class CatalogFileError(ValueError):
    pass


class CatalogValidationError(ValueError):
    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        super().__init__("; ".join(errors))


class CatalogMappingRequiredError(CatalogFileError):
    def __init__(
        self,
        *,
        headers: list[str],
        sample_rows: list[list[str]],
        sheet_name: str | None,
        header_row: int,
    ) -> None:
        super().__init__("Nije moguće automatski prepoznati kolone kataloga.")
        self.headers = headers
        self.sample_rows = sample_rows
        self.sheet_name = sheet_name
        self.header_row = header_row


@dataclass(frozen=True)
class ParsedCatalog:
    checksum: str
    sheet_name: str | None
    products: list[CatalogProductInput]
    aliases_imported: int


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float):
        return str(Decimal(str(value)))
    return str(value).strip()


def _parse_decimal(value: str, field_label: str, row_number: int) -> Decimal | None:
    normalized = value.strip()
    if not normalized:
        return None

    normalized = normalized.replace(" ", "").replace("\u00a0", "")
    if "," in normalized and "." in normalized:
        if normalized.rfind(",") > normalized.rfind("."):
            normalized = normalized.replace(".", "").replace(",", ".")
        else:
            normalized = normalized.replace(",", "")
    elif "," in normalized:
        normalized = normalized.replace(",", ".")

    try:
        result = Decimal(normalized)
    except InvalidOperation as error:
        raise ValueError(f"Red {row_number}: {field_label} nije validan decimalni broj.") from error

    if not result.is_finite():
        raise ValueError(f"Red {row_number}: {field_label} mora biti konačan broj.")
    return result


def _read_csv_rows(file_bytes: bytes) -> list[list[object]]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise CatalogFileError("CSV mora biti kodiran kao UTF-8.") from error

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    return [list(row) for row in csv.reader(io.StringIO(text, newline=""), dialect)]


def _read_xlsx_rows(
    file_bytes: bytes,
    requested_sheet: str | None,
) -> tuple[str, list[list[object]]]:
    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
        raise CatalogFileError("XLSX fajl nije validna Excel radna sveska.")

    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError) as error:
        raise CatalogFileError("XLSX fajl nije moguće otvoriti.") from error

    try:
        if requested_sheet is not None and requested_sheet not in workbook.sheetnames:
            raise CatalogFileError(f"Radni list '{requested_sheet}' ne postoji u XLSX fajlu.")
        worksheets = (
            [workbook[requested_sheet]]
            if requested_sheet is not None
            else list(workbook.worksheets)
        )
        parsed_sheets = [
            (worksheet.title, [list(row) for row in worksheet.iter_rows(values_only=True)])
            for worksheet in worksheets
        ]
        if requested_sheet is not None:
            return parsed_sheets[0]
        ranked_sheets = [
            (candidates[0].score, index, sheet)
            for index, sheet in enumerate(parsed_sheets)
            if (candidates := _catalog_header_candidates(sheet[1]))
        ]
        if ranked_sheets:
            _, _, best_sheet = max(ranked_sheets, key=lambda item: (item[0], -item[1]))
            return best_sheet
        return parsed_sheets[0]
    finally:
        workbook.close()


def _catalog_data_row(row: list[object], columns: dict[str, int]) -> bool:
    values = {
        field: _cell_text(row[index]) if index < len(row) else ""
        for field, index in columns.items()
        if field in {"sku", "name", "unit"}
    }
    return all(values.get(field) for field in ("sku", "name", "unit"))


def _catalog_header_candidates(rows: list[list[object]]) -> list[HeaderCandidate]:
    return find_header_candidates(
        rows,
        aliases=CATALOG_HEADER_ALIASES,
        identity_fields={"sku", "name"},
        minimum_fields=3,
        row_validator=_catalog_data_row,
    )


def _parse_active(value: str, row_number: int) -> bool:
    normalized = normalize_header(value)
    if not normalized:
        return True
    if normalized in ACTIVE_TRUE:
        return True
    if normalized in ACTIVE_FALSE:
        return False
    raise ValueError(f"Red {row_number}: status aktivnog proizvoda nije validan.")


def parse_catalog_file(
    *,
    filename: str,
    file_bytes: bytes,
    mapping: CatalogColumnMapping | None = None,
    sheet_name: str | None = None,
    header_row: int | None = None,
) -> ParsedCatalog:
    suffix = Path(filename.replace("\\", "/")).suffix.lower()
    checksum = hashlib.sha256(file_bytes).hexdigest()

    if suffix == ".csv":
        if b"\x00" in file_bytes[:8192]:
            raise CatalogFileError("CSV mora biti tekstualni fajl.")
        if sheet_name is not None:
            raise CatalogFileError("sheet_name se može koristiti samo za XLSX fajlove.")
        source_rows = _read_csv_rows(file_bytes)
        selected_sheet = None
    elif suffix == ".xlsx":
        selected_sheet, source_rows = _read_xlsx_rows(file_bytes, sheet_name)
    else:
        raise CatalogFileError("Podržani formati kataloga su CSV i XLSX.")

    if not source_rows:
        raise CatalogFileError("Fajl je prazan.")

    header_candidates = _catalog_header_candidates(source_rows)
    if mapping is None:
        if not header_candidates:
            preview_row = max(
                range(min(len(source_rows), 50)),
                key=lambda index: sum(
                    value is not None and bool(str(value).strip()) for value in source_rows[index]
                ),
            )
            raise CatalogMappingRequiredError(
                headers=[_cell_text(value) for value in source_rows[preview_row][:50]],
                sample_rows=[
                    [_cell_text(value) for value in row[:50]]
                    for row in source_rows[preview_row + 1 : preview_row + 6]
                ],
                sheet_name=selected_sheet,
                header_row=preview_row,
            )
        candidate = header_candidates[0]
        try:
            mapping = CatalogColumnMapping.model_validate(candidate.columns)
        except ValidationError as error:
            raise CatalogMappingRequiredError(
                headers=[_cell_text(value) for value in source_rows[candidate.row_index][:50]],
                sample_rows=[
                    [_cell_text(value) for value in row[:50]]
                    for row in source_rows[candidate.row_index + 1 : candidate.row_index + 6]
                ],
                sheet_name=selected_sheet,
                header_row=candidate.row_index,
            ) from error
        header_row = candidate.row_index
    else:
        header_row = (
            header_row
            if header_row is not None
            else header_candidates[0].row_index
            if header_candidates
            else 0
        )

    headers = [_cell_text(value) for value in source_rows[header_row]]
    mapped_columns = [index for index in mapping.model_dump().values() if index is not None]
    if any(index >= len(headers) for index in mapped_columns):
        raise CatalogValidationError(["Mapiranje sadrži kolonu koja ne postoji u zaglavlju fajla."])
    if any(not headers[index] for index in mapped_columns):
        raise CatalogValidationError(["Izabrana kolona nema naziv u zaglavlju fajla."])

    products: list[CatalogProductInput] = []
    errors: list[str] = []
    seen_skus: set[str] = set()
    alias_count = 0

    for row_number, values in enumerate(source_rows[header_row + 1 :], start=header_row + 2):
        row = [_cell_text(value) for value in values]
        if not any(row):
            continue
        if len(products) >= MAX_CATALOG_PRODUCTS:
            raise CatalogValidationError(
                [f"Katalog prelazi pilot limit od {MAX_CATALOG_PRODUCTS} proizvoda."]
            )

        def value_for(field: str) -> str:
            index = getattr(mapping, field)
            return row[index] if index is not None and index < len(row) else ""

        sku = value_for("sku")
        name = value_for("name")
        unit = value_for("unit")
        if not sku and not name and not unit:
            continue
        if not sku or not name or not unit:
            errors.append(f"Red {row_number}: SKU, naziv i jedinica mere su obavezni.")
            if len(errors) >= MAX_ROW_ERRORS:
                errors.append("Prikazano je prvih 100 grešaka.")
                break
            continue

        if sku in seen_skus:
            errors.append(f"Red {row_number}: SKU '{sku}' je dupliran u fajlu.")
            if len(errors) >= MAX_ROW_ERRORS:
                errors.append("Prikazano je prvih 100 grešaka.")
                break
            continue

        try:
            price = _parse_decimal(value_for("price"), "cena", row_number)
            vat_rate = _parse_decimal(value_for("vat_rate"), "PDV stopa", row_number)
            active = _parse_active(value_for("active"), row_number)
        except ValueError as error:
            errors.append(str(error))
            if len(errors) >= MAX_ROW_ERRORS:
                errors.append("Prikazano je prvih 100 grešaka.")
                break
            continue

        aliases: list[str] = []
        seen_aliases: set[str] = set()
        for raw_alias in ALIAS_SEPARATOR.split(value_for("aliases")):
            alias = raw_alias.strip()
            normalized_alias = alias.casefold()
            if alias and normalized_alias not in seen_aliases:
                aliases.append(alias)
                seen_aliases.add(normalized_alias)
        try:
            product = CatalogProductInput(
                source_row=row_number,
                sku=sku,
                name=name,
                unit=unit,
                price=price,
                currency=(value_for("currency").upper() or None),
                vat_rate=vat_rate,
                active=active,
                aliases=aliases,
            )
        except ValidationError as error:
            errors.extend(f"Red {row_number}: {issue['msg']}" for issue in error.errors())
            if len(errors) >= MAX_ROW_ERRORS:
                errors.append("Prikazano je prvih 100 grešaka.")
                break
            continue

        products.append(product)
        seen_skus.add(sku)
        alias_count += len(aliases)

    if errors:
        raise CatalogValidationError(errors)
    if not products:
        raise CatalogValidationError(["Nema validnih proizvoda za uvoz."])

    return ParsedCatalog(
        checksum=checksum,
        sheet_name=selected_sheet,
        products=products,
        aliases_imported=alias_count,
    )
